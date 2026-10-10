"""Build the Supplementary Material as PDF + DOCX from its markdown."""
import os, re
from xml.sax.saxutils import escape
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                TableStyle, PageBreak, Preformatted)
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib import colors

FONT_DIR = r"C:/Windows/Fonts"
pdfmetrics.registerFont(TTFont("Times", os.path.join(FONT_DIR, "times.ttf")))
pdfmetrics.registerFont(TTFont("Times-Bold", os.path.join(FONT_DIR, "timesbd.ttf")))
pdfmetrics.registerFont(TTFont("Times-Italic", os.path.join(FONT_DIR, "timesi.ttf")))
pdfmetrics.registerFont(TTFont("Mono", os.path.join(FONT_DIR, "cour.ttf")))
pdfmetrics.registerFontFamily("Times", normal="Times", bold="Times-Bold",
                              italic="Times-Italic")

MD = r"D:/ace/source/MDYOLO_RSL_supplementary.md"
PDF = r"D:/ace/source/MDYOLO_RSL_supplementary.pdf"
DOCX = r"D:/ace/source/MDYOLO_RSL_supplementary.docx"
CONTENT_W = A4[0] - 2 * cm * 3

def st(name, **kw):
    b = dict(fontName="Times", fontSize=10.5, leading=14, alignment=TA_JUSTIFY)
    b.update(kw); return ParagraphStyle(name, **b)

S_TITLE = st("t", fontName="Times-Bold", fontSize=14, leading=18, alignment=TA_CENTER, spaceAfter=8)
S_H2 = st("h2", fontName="Times-Bold", fontSize=12, leading=15, spaceBefore=10, spaceAfter=5)
S_BODY = st("b", spaceAfter=6)
S_CAP = st("cap", fontName="Times-Italic", fontSize=9, leading=11, alignment=TA_CENTER, spaceBefore=2, spaceAfter=8)
S_CODE = ParagraphStyle("code", fontName="Mono", fontSize=8, leading=10, leftIndent=8,
                        backColor=colors.HexColor("#f4f4f4"), borderPadding=4, spaceAfter=8)

def md_inline(t):
    t = escape(t)
    t = re.sub(r'`([^`]+)`', r'<font face="Mono">\1</font>', t)
    t = re.sub(r'\*\*([^*]+)\*\*', r'<b>\1</b>', t)
    t = re.sub(r'(?<!\*)\*([^*]+)\*(?!\*)', r'<i>\1</i>', t)
    return t

with open(MD, encoding="utf-8") as f:
    lines = f.read().replace("\u2212", "-").split("\n")

story = []
def add_table(header, body):
    data = [[Paragraph(md_inline(c), st("th", fontName="Times-Bold", fontSize=8.5, leading=10)) for c in header]]
    for row in body:
        data.append([Paragraph(md_inline(c), st("tc", fontSize=8.5, leading=10)) for c in row])
    w = [CONTENT_W / len(header)] * len(header)
    t = Table(data, colWidths=w, repeatRows=1)
    t.setStyle(TableStyle([
        ("GRID", (0,0), (-1,-1), 0.4, colors.HexColor("#888888")),
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#e8eef5")),
        ("VALIGN", (0,0), (-1,-1), "TOP"),
        ("LEFTPADDING", (0,0), (-1,-1), 4), ("RIGHTPADDING", (0,0), (-1,-1), 4),
        ("TOPPADDING", (0,0), (-1,-1), 3), ("BOTTOMPADDING", (0,0), (-1,-1), 3)]))
    story.append(t); story.append(Spacer(1, 4))

i = 0
while i < len(lines):
    s = lines[i].strip()
    if s == "" or s in ("---", "***"):
        i += 1; continue
    if s.startswith("```"):
        i += 1; buf = []
        while i < len(lines) and not lines[i].strip().startswith("```"):
            buf.append(lines[i]); i += 1
        i += 1
        story.append(Preformatted("\n".join(buf), S_CODE))
        continue
    if (s.startswith("|") and i+1 < len(lines) and lines[i+1].strip().startswith("|")
            and set(lines[i+1].strip().replace(" ", "")) <= set("|-:")):
        tbl = []; j = i
        while j < len(lines) and lines[j].strip().startswith("|"):
            tbl.append(lines[j].strip()); j += 1
        rows = [[c.strip() for c in r.strip("|").split("|")] for r in tbl]
        add_table(rows[0], rows[2:]); i = j
        while i < len(lines) and lines[i].strip() == "":
            i += 1
        if i < len(lines) and lines[i].strip().startswith("*Table"):
            story.append(Paragraph(md_inline(lines[i].strip()), S_CAP)); i += 1
        continue
    if s.startswith("# "):
        story.append(Paragraph(md_inline(s[2:]), S_TITLE)); i += 1; continue
    if s.startswith("## "):
        story.append(Paragraph(md_inline(s[3:]), S_H2)); i += 1; continue
    story.append(Paragraph(md_inline(s), S_BODY)); i += 1

doc = SimpleDocTemplate(PDF, pagesize=A4, leftMargin=3*cm, rightMargin=3*cm,
                        topMargin=3*cm, bottomMargin=3*cm,
                        title="Supplementary Material: Empirical Reproducibility Audit of MD-YOLO",
                        author="Wang Qinmin")
doc.build(story)
print("Saved supplementary PDF ->", PDF)

# ---- DOCX ----
from docx import Document
from docx.shared import Cm as DCm, Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

d = Document()
for sec in d.sections:
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = DCm(3)
n = d.styles["Normal"]; n.font.name = "Times New Roman"; n.font.size = Pt(11)
n._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

def runs(p, text, size=11, mono=False, italic=False, bold=False):
    r = p.add_run(text); r.font.name = "Courier New" if mono else "Times New Roman"
    r.font.size = Pt(size); r.italic = italic; r.bold = bold

i = 0
while i < len(lines):
    s = lines[i].strip()
    if s == "" or s in ("---", "***"):
        i += 1; continue
    if s.startswith("```"):
        i += 1; buf = []
        while i < len(lines) and not lines[i].strip().startswith("```"):
            buf.append(lines[i]); i += 1
        i += 1
        p = d.add_paragraph(); p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_after = Pt(8)
        runs(p, "\n".join(buf), size=9, mono=True)
        continue
    if (s.startswith("|") and i+1 < len(lines) and lines[i+1].strip().startswith("|")
            and set(lines[i+1].strip().replace(" ", "")) <= set("|-:")):
        tbl = []; j = i
        while j < len(lines) and lines[j].strip().startswith("|"):
            tbl.append(lines[j].strip()); j += 1
        rows = [[c.strip() for c in r.strip("|").split("|")] for r in tbl]
        t = d.add_table(rows=1, cols=len(rows[0])); t.style = "Table Grid"
        for jj, h in enumerate(rows[0]):
            runs(t.rows[0].cells[jj].paragraphs[0], h, size=9, bold=True)
        for row in rows[2:]:
            cells = t.add_row().cells
            for jj, v in enumerate(row):
                runs(cells[jj].paragraphs[0], v, size=9)
        d.add_paragraph()
        i = j
        while i < len(lines) and lines[i].strip() == "":
            i += 1
        if i < len(lines) and lines[i].strip().startswith("*Table"):
            cap = d.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            runs(cap, lines[i].strip().strip("*"), size=9, italic=True); i += 1
        continue
    if s.startswith("# "):
        p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        runs(p, s[2:], size=14, bold=True); i += 1; continue
    if s.startswith("## "):
        p = d.add_paragraph(); runs(p, s[3:], size=12, bold=True); i += 1; continue
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    runs(p, s, size=11); i += 1

try:
    d.save(DOCX); print("Saved supplementary DOCX ->", DOCX)
except PermissionError:
    alt = r"D:/ace/source/MDYOLO_RSL_supplementary_submitted.docx"
    d.save(alt); print("Saved supplementary DOCX ->", alt)
