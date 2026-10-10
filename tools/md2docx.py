"""
将论文 Markdown 转换为排版良好的 .docx (学术风格)。
支持: 标题层级、加粗/行内代码、GFM 表格、无序列表、代码块、普通段落。
"""
import re, sys
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def add_inline(paragraph, text):
    """处理 **bold** 与 `code` 行内格式。"""
    # 先按 code 拆分, 再在每段里处理 bold
    parts = re.split(r'(`[^`]+`)', text)
    for p in parts:
        if not p:
            continue
        if p.startswith('`') and p.endswith('`') and len(p) > 1:
            r = paragraph.add_run(p[1:-1])
            r.font.name = 'Consolas'
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0x8B, 0x00, 0x00)
        else:
            # bold 处理
            bparts = re.split(r'(\*\*[^*]+\*\*)', p)
            for b in bparts:
                if not b:
                    continue
                if b.startswith('**') and b.endswith('**') and len(b) > 2:
                    r = paragraph.add_run(b[2:-2])
                    r.bold = True
                else:
                    paragraph.add_run(b)

def convert(md_path, docx_path):
    doc = Document()
    # 基础样式
    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    with open(md_path, encoding='utf-8') as f:
        lines = f.read().split('\n')

    i = 0
    in_code = False
    code_buf = []
    while i < len(lines):
        line = lines[i]
        if line.strip().startswith('```'):
            if not in_code:
                in_code = True
                code_buf = []
            else:
                in_code = False
                # 输出代码块
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.3)
                run = p.add_run('\n'.join(code_buf))
                run.font.name = 'Consolas'
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
                code_buf = []
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue

        s = line.strip()
        if not s:
            i += 1
            continue

        # 表格 (兼容 `|---|` 与 `| --- |` 两种分隔行写法)
        if s.startswith('|') and i + 1 < len(lines):
            nxt = lines[i+1].strip()
            if nxt.startswith('|') and '-' in nxt and set(nxt.replace(' ', '')) <= set('|-:'):
                tbl_lines = []
            j = i
            while j < len(lines) and lines[j].strip().startswith('|'):
                tbl_lines.append(lines[j].strip())
                j += 1
            rows = []
            for tl in tbl_lines:
                cells = [c.strip() for c in tl.strip('|').split('|')]
                rows.append(cells)
            header = rows[0]
            body = rows[2:]  # 跳过分隔行
            table = doc.add_table(rows=1, cols=len(header))
            table.style = 'Light Grid Accent 1'
            hdr = table.rows[0].cells
            for k, h in enumerate(header):
                hdr[k].text = ''
                add_inline(hdr[k].paragraphs[0], h)
                for rr in hdr[k].paragraphs[0].runs:
                    rr.bold = True
            for row in body:
                cells = table.add_row().cells
                for k, c in enumerate(row):
                    cells[k].text = ''
                    add_inline(cells[k].paragraphs[0], c)
            doc.add_paragraph()
            i = j
            continue

        # 标题
        if s.startswith('### '):
            p = doc.add_heading(s[4:], level=2)
            i += 1
            continue
        if s.startswith('## '):
            p = doc.add_heading(s[3:], level=1)
            i += 1
            continue
        if s.startswith('# '):
            p = doc.add_heading(s[2:], level=0)
            i += 1
            continue

        # 列表
        if s.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            add_inline(p, s[2:])
            i += 1
            continue

        # 普通段落
        p = doc.add_paragraph()
        add_inline(p, s)
        i += 1

    doc.save(docx_path)
    print(f"Saved -> {docx_path}")

if __name__ == '__main__':
    convert(sys.argv[1], sys.argv[2])
