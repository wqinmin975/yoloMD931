"""Build the Remote Sensing Letters cover letter as PDF and DOCX."""
import os
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.lib.enums import TA_JUSTIFY, TA_LEFT
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

FONT_DIR = r"C:/Windows/Fonts"
pdfmetrics.registerFont(TTFont("Times", os.path.join(FONT_DIR, "times.ttf")))
pdfmetrics.registerFont(TTFont("Times-Bold", os.path.join(FONT_DIR, "timesbd.ttf")))
pdfmetrics.registerFont(TTFont("Times-Italic", os.path.join(FONT_DIR, "timesi.ttf")))
pdfmetrics.registerFontFamily("Times", normal="Times", bold="Times-Bold",
                              italic="Times-Italic")

PDF = r"D:/ace/source/RSL_cover_letter.pdf"
DOCX = r"D:/ace/source/RSL_cover_letter.docx"

BODY = [
    ("p", "Date: [Date]"),
    ("sp", ""),
    ("p", "To the Editorial Office, Remote Sensing Letters"),
    ("sp", ""),
    ("p", "Dear Editor,"),
    ("sp", ""),
    ("p", "We are pleased to submit our manuscript entitled \u201cAn Empirical "
          "Reproducibility Audit of MD-YOLO for Small-Object Detection in "
          "UAV Imagery\u201d "
          "for consideration as a Research Letter in Remote Sensing Letters."),
    ("p", "The manuscript reports an end-to-end empirical validation of the "
          "MD-YOLO detector for UAV aerial small-object detection on the real "
          "VisDrone2019-DET dataset (7 019 images, 381 963 bounding boxes). "
          "Executing the pipeline on real data\u2014rather than relying on static "
          "source inspection\u2014exposed four runtime defects, the most severe of "
          "which silently produced mAP = 0. After correction, we decompose the "
          "residual performance of the model into a resolved localization "
          "deficit and an open classification deficit attributable to a limited "
          "data\u2013compute regime, and we further quantify a severe zero-shot "
          "domain shift on UAVDT-M (mean best-IoU 0.40 \u2192 0.271). The work "
          "deliberately isolates three non-conflating failure categories and "
          "reports the negative results honestly, which we believe is of direct "
          "value to reproducible remote-sensing research."),
    ("p", "We consider the contribution well suited to Remote Sensing Letters "
          "because it addresses a core UAV remote-sensing task\u2014small-object "
          "detection in aerial imagery\u2014and combines a cross-dataset domain-shift "
          "analysis (VisDrone \u2192 UAVDT-M) with a reproducibility methodology. The "
          "compact Letter format matches the focused scope of this study."),
    ("sp", ""),
    ("p", "Authors: Qinmin Wang, Zhenzhen Hu, Hongyuan Wang, and Yuming Sun "
          "(corresponding author)."),
    ("p", "Manuscript type: Research Letter. Main text is within the 10-page "
          "limit; 1 figure and 3 tables (plus 1 supplementary table)."),
    ("sp", ""),
    ("p", "We confirm the following:"),
    ("b", "\u2022  This manuscript is original, has not been published previously, and is not under consideration by any other journal."),
    ("b", "\u2022  We request subscription-based (non-open-access) publication; we do not wish to publish under an open-access licence and will not pay an article-processing charge."),
    ("b", "\u2022  Full reproducibility scripts, pseudo-code, and supplementary diagnostic tables are supplied as Supplementary Material."),
    ("b", "\u2022  All authors have approved the submission and declare no competing interests."),
    ("sp", ""),
    ("p", "Thank you for your time and consideration. We look forward to your "
          "response."),
    ("sp", ""),
    ("p", "Sincerely,"),
    ("sp", ""),
    ("p", "Yuming Sun"),
    ("p", "School of Computer Science and Artificial Intelligence, Guangdong University of Science and Technology, Dongguan, China"),
    ("p", "Email: sunyuming@gdust.edu.cn"),
    ("p", "ORCID: 0009-0003-7222-2390"),
]

S_P = ParagraphStyle("p", fontName="Times", fontSize=11, leading=15,
                     alignment=TA_JUSTIFY, spaceAfter=6)
S_B = ParagraphStyle("b", fontName="Times", fontSize=11, leading=15,
                     alignment=TA_LEFT, spaceAfter=3, leftIndent=14)
S_N = ParagraphStyle("n", fontName="Times", fontSize=11, leading=15,
                     alignment=TA_LEFT, spaceAfter=4)

story = []
for kind, text in BODY:
    if kind == "sp":
        story.append(Spacer(1, 6))
    elif kind == "b":
        story.append(Paragraph(text, S_B))
    else:
        story.append(Paragraph(text, S_P))

doc = SimpleDocTemplate(PDF, pagesize=A4,
                        leftMargin=2.5*cm, rightMargin=2.5*cm,
                        topMargin=2.5*cm, bottomMargin=2.5*cm,
                        title="Cover Letter - Remote Sensing Letters",
                        author="Yuming Sun")
doc.build(story)
print("Saved cover letter PDF ->", PDF)

# ---- DOCX version ----
from docx import Document
from docx.shared import Cm as DCm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

d = Document()
for sec in d.sections:
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = DCm(2.5)
n = d.styles["Normal"]
n.font.name = "Times New Roman"; n.font.size = Pt(11)
n._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
for kind, text in BODY:
    if kind == "sp":
        continue
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if kind == "p" else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(6 if kind == "p" else 3)
    r = p.add_run(text)
    r.font.name = "Times New Roman"; r.font.size = Pt(11)
try:
    d.save(DOCX)
    print("Saved cover letter DOCX ->", DOCX)
except PermissionError:
    alt = r"D:/ace/source/RSL_cover_letter_submitted.docx"
    d.save(alt); print("Saved cover letter DOCX ->", alt)
